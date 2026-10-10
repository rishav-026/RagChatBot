import io
from typing import List, Optional
from pydantic import BaseModel
import pypdf
import docx

from app.rag.ingestion.cleaner import clean_text

class PageContent(BaseModel):
    page_number: Optional[int] = None
    text: str

class ExtractedDocument(BaseModel):
    filename: str
    file_type: str
    pages: List[PageContent]
    full_text: str

class DocumentExtractionError(Exception):
    """Exception raised when document parsing fails."""
    pass

class UnsupportedFileTypeError(Exception):
    """Exception raised when file format is not supported."""
    pass

class EmptyDocumentError(Exception):
    """Exception raised when extracted document contains no usable text."""
    pass


class BaseExtractor:
    def extract(self, content: bytes, filename: str) -> ExtractedDocument:
        raise NotImplementedError


class PDFExtractor(BaseExtractor):
    def extract(self, content: bytes, filename: str) -> ExtractedDocument:
        try:
            reader = pypdf.PdfReader(io.BytesIO(content))
            pages: List[PageContent] = []
            full_text_parts: List[str] = []

            for idx, page in enumerate(reader.pages):
                raw_text = page.extract_text() or ""
                cleaned = clean_text(raw_text)
                page_num = idx + 1
                if cleaned:
                    pages.append(PageContent(page_number=page_num, text=cleaned))
                    full_text_parts.append(cleaned)

            full_text = "\n\n".join(full_text_parts).strip()
            if not full_text:
                raise EmptyDocumentError(f"PDF document '{filename}' contains no extractable text.")

            return ExtractedDocument(
                filename=filename,
                file_type="pdf",
                pages=pages,
                full_text=full_text
            )
        except EmptyDocumentError:
            raise
        except Exception as e:
            raise DocumentExtractionError(f"Failed to extract text from PDF '{filename}': {str(e)}")


class TXTExtractor(BaseExtractor):
    def extract(self, content: bytes, filename: str) -> ExtractedDocument:
        try:
            # Try UTF-8 first, fallback to latin-1
            try:
                raw_text = content.decode("utf-8")
            except UnicodeDecodeError:
                raw_text = content.decode("latin-1")

            cleaned = clean_text(raw_text)
            if not cleaned:
                raise EmptyDocumentError(f"TXT document '{filename}' is empty.")

            return ExtractedDocument(
                filename=filename,
                file_type="txt",
                pages=[PageContent(page_number=1, text=cleaned)],
                full_text=cleaned
            )
        except EmptyDocumentError:
            raise
        except Exception as e:
            raise DocumentExtractionError(f"Failed to read text file '{filename}': {str(e)}")


class DOCXExtractor(BaseExtractor):
    def extract(self, content: bytes, filename: str) -> ExtractedDocument:
        try:
            doc = docx.Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            # Extract text from tables if available
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)

            raw_text = "\n\n".join(paragraphs)
            cleaned = clean_text(raw_text)

            if not cleaned:
                raise EmptyDocumentError(f"DOCX document '{filename}' contains no extractable text.")

            return ExtractedDocument(
                filename=filename,
                file_type="docx",
                pages=[PageContent(page_number=1, text=cleaned)],
                full_text=cleaned
            )
        except EmptyDocumentError:
            raise
        except Exception as e:
            raise DocumentExtractionError(f"Failed to extract text from DOCX '{filename}': {str(e)}")


class DocumentExtractorFactory:
    _extractors = {
        "pdf": PDFExtractor(),
        "txt": TXTExtractor(),
        "docx": DOCXExtractor()
    }

    @classmethod
    def get_extractor(cls, filename: str) -> BaseExtractor:
        ext = filename.split(".")[-1].lower() if "." in filename else ""
        if ext not in cls._extractors:
            raise UnsupportedFileTypeError(
                f"Unsupported file format '.{ext}'. Supported formats: PDF, TXT, DOCX."
            )
        return cls._extractors[ext]

    @classmethod
    def extract(cls, content: bytes, filename: str) -> ExtractedDocument:
        extractor = cls.get_extractor(filename)
        return extractor.extract(content, filename)
